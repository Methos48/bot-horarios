import asyncio
import base64
import io
import os
import threading
import time
from bs4 import BeautifulSoup
import discord
from flask import Flask
import openpyxl
from PIL import Image, ImageDraw, ImageFont
import requests

# --- CONFIGURACIÓN GENERAL ---
app = Flask(__name__)
DISCORD_TOKEN = os.getenv("DISCORD_TOKEN")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

DISCORD_CANAL_HORARIOS_ID = 1548528724268552263
DISCORD_CANAL_RONDA_ID = 1548528618949582929

EXCEL_PATH = os.getenv("EXCEL_PATH", "Calculadora de RAID.xlsm")
EXCEL_URL = "https://1drv.ms/x/c/434ba5d6d0d889c3/IQAwNBrAH5eLQZWV-N3ufOfYAY8sBOApZdxzU8GuWMEBs0E?download=1"

intents = discord.Intents.default()
intents.message_content = True
client_discord = discord.Client(intents=intents)

ultimo_mensaje_horarios_id = None
ultimo_mensaje_ronda_id = None


@app.route("/")
def home():
    return "🤖 Bot autónomo de Lineage II activo 24/7!"


def run_web():
    port = int(os.getenv("PORT", 8080))
    app.run(host="0.0.0.0", port=port)


def descargar_excel_desde_onedrive():
    print("☁️ Descargando la última versión del Excel desde OneDrive...")
    try:
        response = requests.get(EXCEL_URL, timeout=60)
        if response.status_code == 200:
            with open(EXCEL_PATH, "wb") as f:
                f.write(response.content)
            print("✅ ¡Excel descargado y actualizado con éxito desde OneDrive!")
            return True
        else:
            print(f"❌ Error al descargar de OneDrive. Código HTTP: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Error de red al intentar descargar el Excel: {e}")
        return False


def cargar_excel():
    try:
        if not os.path.exists(EXCEL_PATH):
            print("⚠️ El archivo no está localmente, intentando descargar de OneDrive...")
            descargar_excel_desde_onedrive()
            
        if os.path.exists(EXCEL_PATH):
            # keep_vba=True garantiza que las macros no se pierdan al leer/escribir
            return openpyxl.load_workbook(EXCEL_PATH, keep_vba=True)
        else:
            print(f"❌ No se encontró el archivo Excel en la ruta: {EXCEL_PATH}")
            return None
    except Exception as e:
        print(f"❌ Error al abrir el Excel: {e}")
        return None


def guardar_excel(wb):
    try:
        wb.save(EXCEL_PATH)
        return True
    except Exception as e:
        print(f"❌ Error al guardar el Excel: {e}")
        return False


def generar_imagen_horario_rojo(wb):
    try:
        img_width, img_height = 800, 900
        img = Image.new("RGB", (img_width, img_height), color="#FDF3D8")
        draw = ImageDraw.Draw(img)

        try:
            font_titulo = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 22
            )
            font_texto = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 16
            )
        except:
            font_titulo = ImageFont.load_default()
            font_texto = ImageFont.load_default()

        draw.rectangle([0, 0, img_width, 110], fill="#8B0000")
        draw.text(
            (250, 40),
            "OKT Raid OKT Ally HTF",
            fill="#FFD700",
            font=font_titulo,
        )

        draw.text((50, 140), "RAID", fill="#003366", font=font_texto)
        draw.text((200, 140), "DIA", fill="#003366", font=font_texto)
        draw.text((310, 140), "FECHA", fill="#003366", font=font_texto)
        draw.text((530, 140), "ARG / CHI", fill="#006600", font=font_texto)
        draw.text((640, 140), "VEN", fill="#006600", font=font_texto)
        draw.text((720, 140), "ESP", fill="#006600", font=font_texto)

        draw.line([30, 175, 770, 175], fill="#C08040", width=2)

        sheet = (
            wb["HORARIO ROJO"]
            if "HORARIO ROJO" in wb.sheetnames
            else wb["CALCULADORA"]
        )
        y_offset = 200
        for row in range(10, 25):
            raid_nombre = sheet.cell(row=row, column=1).value
            if not raid_nombre:
                break

            dia = str(sheet.cell(row=row, column=2).value or "")
            fecha = str(sheet.cell(row=row, column=3).value or "")
            arg = str(sheet.cell(row=row, column=6).value or "18:00")
            ven = str(sheet.cell(row=row, column=7).value or "17:00")
            esp = str(sheet.cell(row=row, column=8).value or "23:00")

            color_texto = "#CC0000" if row in [13, 19, 22, 23] else "#003300"

            draw.text(
                (50, y_offset), str(raid_nombre), fill=color_texto, font=font_texto
            )
            draw.text((200, y_offset), dia, fill=color_texto, font=font_texto)
            draw.text((310, y_offset), fecha, fill=color_texto, font=font_texto)
            draw.text((530, y_offset), arg, fill=color_texto, font=font_texto)
            draw.text((640, y_offset), ven, fill=color_texto, font=font_texto)
            draw.text((720, y_offset), esp, fill=color_texto, font=font_texto)

            y_offset += 35

        output_img = io.BytesIO()
        img.save(output_img, format="PNG")
        output_img.seek(0)
        return output_img
    except Exception as e:
        print(f"❌ Error generando Horario Rojo: {e}")
        return None


def generar_imagen_ronda_rojo(wb):
    try:
        img_width, img_height = 850, 950
        img = Image.new("RGB", (img_width, img_height), color="#FDF3D8")
        draw = ImageDraw.Draw(img)

        try:
            font_titulo = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 22
            )
            font_texto = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 13
            )
        except:
            font_titulo = ImageFont.load_default()
            font_texto = ImageFont.load_default()

        draw.rectangle([0, 0, img_width, 110], fill="#8B0000")
        draw.text(
            (250, 40),
            "OKT Ronda Raid HTF",
            fill="#FFD700",
            font=font_titulo,
        )

        draw.text((40, 135), "Raid", fill="#003366", font=font_texto)
        draw.text((260, 135), "LVL", fill="#003366", font=font_texto)
        draw.text((310, 135), "Hora", fill="#003366", font=font_texto)

        draw.text((450, 135), "Raid", fill="#003366", font=font_texto)
        draw.text((670, 135), "LVL", fill="#003366", font=font_texto)
        draw.text((720, 135), "Hora", fill="#003366", font=font_texto)

        draw.line([30, 160, 820, 160], fill="#C08040", width=2)

        sheet = (
            wb["RONDA ROJO"] if "RONDA ROJO" in wb.sheetnames else wb["CALCULADORA"]
        )

        y_left = 180
        for row in range(9, 30):
            raid = sheet.cell(row=row, column=2).value
            if not raid:
                break
            lvl = str(sheet.cell(row=row, column=4).value or "")
            hora = str(sheet.cell(row=row, column=5).value or "")

            draw.text((40, y_left), str(raid), fill="#003300", font=font_texto)
            draw.text((260, y_left), lvl, fill="#003300", font=font_texto)
            draw.text((310, y_left), hora, fill="#006600", font=font_texto)
            y_left += 28

        y_right = 180
        for row in range(9, 35):
            raid = sheet.cell(row=row, column=7).value
            if not raid:
                break
            lvl = str(sheet.cell(row=row, column=9).value or "")
            hora = str(sheet.cell(row=row, column=10).value or "")

            draw.text((450, y_right), str(raid), fill="#003300", font=font_texto)
            draw.text((670, y_right), lvl, fill="#003300", font=font_texto)
            draw.text((720, y_right), hora, fill="#006600", font=font_texto)
            y_right += 28

        output_img = io.BytesIO()
        img.save(output_img, format="PNG")
        output_img.seek(0)
        return output_img
    except Exception as e:
        print(f"❌ Error generando Ronda Rojo: {e}")
        return None


@client_discord.event
async def on_ready():
    print(f"🤖 Bot conectado exitosamente como {client_discord.user}")


def extraer_datos_imagen(img_pil):
    buffered = io.BytesIO()
    img_pil.save(buffered, format="JPEG")
    img_base64 = base64.b64encode(buffered.getvalue()).decode("utf-8")

    url = f"https://generativelanguage.googleapis.com/v1/models/gemini-3.6-flash:generateContent?key={GEMINI_API_KEY}"
    headers = {"Content-Type": "application/json"}

    prompt_estricto = (
        "Extrae los nombres de los raids y sus horarios de esta imagen."
        " Devuelve UNICAMENTE lineas con el formato exacto 'Nombre: Horario'."
        " Prohibido usar saludos, explicaciones, markdown, asteriscos o tablas."
    )

    payload = {
        "contents": [{
            "parts": [
                {
                    "inline_data": {
                        "mime_type": "image/jpeg",
                        "data": img_base64,
                    }
                },
                {"text": prompt_estricto},
            ]
        }]
    }

    intentos = 3
    for i in range(intentos):
        try:
            response = requests.post(url, headers=headers, json=payload, timeout=30)
            if response.status_code == 200:
                data = response.json()
                texto_resultado = (
                    data.get("candidates", [{}])[0]
                    .get("content", {})
                    .get("parts", [{}])[0]
                    .get("text", "")
                )
                return texto_resultado.strip().split("\n")
            else:
                raise Exception(f"HTTP {response.status_code}: {response.text}")
        except Exception as ex:
            print(f"⚠️ Intento {i+1} fallido por API REST: {ex}")
            if i < intentos - 1:
                time.sleep(4)
            else:
                raise ex


@client_discord.event
async def on_message(message):
    global ultimo_mensaje_horarios_id, ultimo_mensaje_ronda_id

    if message.author == client_discord.user:
        return

    if message.attachments:
        imagenes_validas = [
            att
            for att in message.attachments
            if att.filename.lower().endswith((".png", ".jpg", ".jpeg", ".webp"))
        ]

        if imagenes_validas:
            print(
                f"📸 Mensaje detectado con {len(imagenes_validas)} imagen(es) de"
                " raids."
            )
            diccionario_raids_consolidado = {}

            mapa_raids = {
                "valakas": "Valakas",
                "balrog": "Balrog",
                "barakiel": "Barakiel",
                "core": "Core",
                "orfen": "Orfen",
                "antharas": "Antharas",
                "electrical": "Electrical",
                "baium": "Baium",
                "zaken": "Zaken",
                "frintezza": "Frintezza",
                "fafureon": "Fafureon",
                "queen ant": "Queen Ant",
                "freya": "Freya",
                "zariche": "Zariche",
            }

            raids_oficiales = list(mapa_raids.values())

            bytes_imagenes = []
            for attachment in imagenes_validas:
                try:
                    b = await attachment.read()
                    bytes_imagenes.append(b)
                except Exception as e:
                    print(f"❌ Error al leer el adjunto {attachment.filename}: {e}")

            for idx_img, img_bytes in enumerate(bytes_imagenes):
                try:
                    img_pil = Image.open(io.BytesIO(img_bytes))
                    lineas_extraidas = await asyncio.to_thread(
                        extraer_datos_imagen, img_pil
                    )

                    if lineas_extraidas:
                        for linea in lineas_extraidas:
                            linea_limpia = (
                                linea.replace("|", "")
                                .replace("*", "")
                                .replace("`", "")
                                .replace("-", "")
                                .strip()
                            )
                            if not linea_limpia or ":" not in linea_limpia:
                                continue

                            partes = linea_limpia.split(":", 1)
                            nombre_leido = partes[0].strip().lower()
                            horario = partes[1].strip()

                            nombre_encontrado = None
                            for clave, oficial in mapa_raids.items():
                                if clave in nombre_leido:
                                    nombre_encontrado = oficial
                                    break

                            if nombre_encontrado and horario:
                                diccionario_raids_consolidado[nombre_encontrado] = horario

                    if idx_img < len(bytes_imagenes) - 1:
                        time.sleep(3)

                except Exception as ex:
                    print(f"⚠️ Error procesando imagen en memoria: {ex}")

            if len(diccionario_raids_consolidado) > 0:
                try:
                    wb = cargar_excel()
                    if wb and "CALCULADORA" in wb.sheetnames:
                        sheet = wb["CALCULADORA"]

                        print(
                            "✍️ Rellenando la columna B (B2:B15) con los datos"
                            " consolidados..."
                        )
                        for idx, raid_oficial in enumerate(raids_oficiales):
                            fila = idx + 2
                            horario_valor = diccionario_raids_consolidado.get(raid_oficial, "")
                            sheet.cell(row=fila, column=2, value=horario_valor)

                        guardar_excel(wb)

                    wb_verificacion = cargar_excel()
                    casillas_llenas = True
                    if wb_verificacion and "CALCULADORA" in wb_verificacion.sheetnames:
                        sheet_v = wb_verificacion["CALCULADORA"]
                        for row in range(2, 16):
                            val = sheet_v.cell(row=row, column=2).value
                            if not val or str(val).strip() == "":
                                casillas_llenas = False
                                break

                    if wb_verificacion:
                        img_horarios = generar_imagen_horario_rojo(wb_verificacion)
                        canal_horarios = client_discord.get_channel(
                            int(DISCORD_CANAL_HORARIOS_ID)
                        )
                        if canal_horarios and img_horarios:
                            if ultimo_mensaje_horarios_id:
                                try:
                                    msg_ant = await canal_horarios.fetch_message(
                                        ultimo_mensaje_horarios_id
                                    )
                                    await msg_ant.delete()
                                except:
                                    pass
                            f_horarios = discord.File(
                                fp=img_horarios, filename="Horario_Rojo.png"
                            )
                            msg_h = await canal_horarios.send(
                                "🔥 **HORARIOS DE RAIDS ACTUALIZADOS (vía imagen):**",
                                file=f_horarios,
                            )
                            ultimo_mensaje_horarios_id = msg_h.id

                        img_ronda = generar_imagen_ronda_rojo(wb_verificacion)
                        canal_ronda = client_discord.get_channel(
                            int(DISCORD_CANAL_RONDA_ID)
                        )
                        if canal_ronda and img_ronda:
                            if ultimo_mensaje_ronda_id:
                                try:
                                    msg_ant_r = await canal_ronda.fetch_message(
                                        ultimo_mensaje_ronda_id
                                    )
                                    await msg_ant_r.delete()
                                except:
                                    pass
                            f_ronda = discord.File(fp=img_ronda, filename="Ronda_Rojo.png")
                            msg_r = await canal_ronda.send(
                                "⚔️ **RONDA ROJO (Nivel 60+) ACTUALIZADA (vía"
                                " imagen):**",
                                file=f_ronda,
                            )
                            ultimo_mensaje_ronda_id = msg_r.id

                    if casillas_llenas:
                        print(
                            "✔️ Verificación exitosa: Rango B2:B15 completo. Borrando"
                            " mensaje original de las fotos..."
                        )
                        try:
                            await message.delete()
                        except Exception as e:
                            print(f"❌ No se pudo borrar el mensaje original: {e}")
                    else:
                        print(
                            "⚠️ Advertencia: Faltan celdas por completar en B2:B15, el"
                            " mensaje original no se borrará."
                        )

                except Exception as e:
                    print(f"❌ Error general procesando el Excel y las imágenes: {e}")


if __name__ == "__main__":
    # Descargar el archivo directamente de OneDrive al arrancar
    descargar_excel_desde_onedrive()

    t = threading.Thread(target=run_web)
    t.daemon = True
    t.start()

    client_discord.run(DISCORD_TOKEN)
